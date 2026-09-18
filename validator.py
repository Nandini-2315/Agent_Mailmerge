import re
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional, Set
import docx

from .logger import logger
from .excel_manager import load_config


class ValidationError(Exception):
    """Raised when validation fails at the dataset or configuration level."""
    pass


def validate_environment(
    excel_path: Optional[Path] = None,
    template_path: Optional[Path] = None,
    output_dir: Optional[Path] = None,
    config: Optional[Dict[str, Any]] = None
) -> Tuple[Path, Path, Path]:
    """
    Validate that required files and directories exist.
    Resolves fallback template paths if primary template is missing.
    Returns (resolved_excel_path, resolved_template_path, resolved_output_dir).
    """
    cfg = config or load_config()
    
    # 1. Validate Excel Path
    e_path = Path(excel_path or cfg["paths"]["excel_file"])
    if not e_path.exists():
        raise FileNotFoundError(f"Excel file not found at: {e_path.resolve()}")

    # 2. Validate Template Path (with fallback support)
    t_path = Path(template_path or cfg["paths"]["template_file"])
    if not t_path.exists():
        fb_path = Path(cfg["paths"].get("fallback_template_file", "input/Student_Certificate.docx"))
        if fb_path.exists():
            logger.info(f"Template '{t_path}' not found. Using fallback template '{fb_path}'.")
            t_path = fb_path
        else:
            raise FileNotFoundError(
                f"Certificate template not found at '{t_path.resolve()}' and fallback '{fb_path.resolve()}' also does not exist."
            )

    # 3. Validate / Create Output Directory
    o_dir = Path(output_dir or cfg["paths"]["output_dir"])
    o_dir.mkdir(parents=True, exist_ok=True)

    return e_path, t_path, o_dir


def validate_dataset_structure(students: List[Dict[str, Any]], config: Optional[Dict[str, Any]] = None) -> None:
    """
    Validates dataset-level integrity:
    1. Check for empty dataset.
    2. Check for duplicate Certificate_IDs. Duplicate IDs halt the process as required by Section 14.
    3. Check that required columns exist in records.
    """
    cfg = config or load_config()
    required_cols = cfg["columns"].get("required_columns", ["Certificate_ID", "Name"])
    id_col = cfg["columns"].get("id_column", "Certificate_ID")

    if not students:
        raise ValidationError("The Excel file contains no student records.")

    # Check required columns present
    sample = students[0]
    missing_cols = [col for col in required_cols if col not in sample]
    if missing_cols:
        raise ValidationError(
            f"Excel file is missing required column(s): {', '.join(missing_cols)}. "
            f"Available columns: {', '.join([k for k in sample.keys() if not k.startswith('_')])}"
        )

    # Check for duplicate Certificate_IDs
    id_counts: Dict[str, List[int]] = {}
    for student in students:
        cert_id = str(student.get(id_col, "")).strip()
        row_num = student.get("_row_number", 0)
        if cert_id:
            id_counts.setdefault(cert_id, []).append(row_num)

    duplicates = {cid: rows for cid, rows in id_counts.items() if len(rows) > 1}
    if duplicates:
        err_msg_parts = []
        for cid, rows in duplicates.items():
            err_msg_parts.append(f"'{cid}' appears {len(rows)} times (Excel rows: {rows})")
        full_err = "Duplicate Certificate_ID found: " + "; ".join(err_msg_parts) + ". Each student must have a unique Certificate_ID."
        logger.error(full_err)
        raise ValidationError(full_err)


def validate_students(
    students: List[Dict[str, Any]],
    config: Optional[Dict[str, Any]] = None
) -> Tuple[List[Dict[str, Any]], Dict[str, str]]:
    """
    Validates individual student records.
    Returns:
        (valid_students, failed_records_dict)
    where failed_records_dict maps identifier -> failure reason.
    """
    cfg = config or load_config()
    required_cols = cfg["columns"].get("required_columns", ["Certificate_ID", "Name"])
    id_col = cfg["columns"].get("id_column", "Certificate_ID")

    valid_students: List[Dict[str, Any]] = []
    failed_students: Dict[str, str] = {}

    for idx, student in enumerate(students):
        cert_id = str(student.get(id_col, "")).strip()
        row_num = student.get("_row_number", idx + 2)
        identifier = cert_id or f"Row_{row_num}"

        if not cert_id:
            failed_students[identifier] = f"Missing required '{id_col}' at Excel row {row_num}."
            logger.warning(f"Row {row_num} rejected: {failed_students[identifier]}")
            continue

        missing_fields = []
        for req in required_cols:
            val = str(student.get(req, "")).strip()
            if not val:
                missing_fields.append(req)

        if missing_fields:
            reason = f"Missing required field(s): {', '.join(missing_fields)}"
            failed_students[identifier] = reason
            logger.warning(f"{identifier} rejected: {reason}")
            continue

        valid_students.append(student)

    return valid_students, failed_students


def extract_template_placeholders(template_path: Path) -> Set[str]:
    """
    Extract all placeholder tags in the form of {{field}} from the DOCX template.
    Checks body paragraphs, tables, drawings, headers, and footers.
    """
    try:
        doc = docx.Document(template_path)
    except Exception as e:
        raise ValidationError(f"Failed to open template DOCX at {template_path}: {e}")

    placeholders: Set[str] = set()
    pattern = re.compile(r"\{\{([a-zA-Z0-9_\s\-]+)\}\}")

    # Inspect all w:p elements across entire document
    for p_elem in doc._element.xpath(".//w:p"):
        p = docx.text.paragraph.Paragraph(p_elem, doc)
        full_text = "".join(r.text for r in p.runs)
        for match in pattern.findall(full_text):
            placeholders.add(match.strip())

    # Inspect sections (headers / footers)
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header and not header.is_linked_to_previous:
                for p_elem in header._element.xpath(".//w:p"):
                    p = docx.text.paragraph.Paragraph(p_elem, doc)
                    full_text = "".join(r.text for r in p.runs)
                    for match in pattern.findall(full_text):
                        placeholders.add(match.strip())

        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            if footer and not footer.is_linked_to_previous:
                for p_elem in footer._element.xpath(".//w:p"):
                    p = docx.text.paragraph.Paragraph(p_elem, doc)
                    full_text = "".join(r.text for r in p.runs)
                    for match in pattern.findall(full_text):
                        placeholders.add(match.strip())

    return placeholders
