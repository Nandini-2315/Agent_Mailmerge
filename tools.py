from pathlib import Path
import pandas as pd
from docx import Document


BASE_DIR = Path(__file__).resolve().parent.parent

EXCEL_FILE = BASE_DIR / "input" / "studentsdata.xlsx"
TEMPLATE_FILE = BASE_DIR / "input" / "certificatetemp.docx"
OUTPUT_DIR = BASE_DIR / "output"


def read_student_data():
    """
    Read student information from the Excel file.

    Returns:
        list: A list of dictionaries containing student information.
    """

    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    df = pd.read_excel(EXCEL_FILE)

    if df.empty:
        raise ValueError("The Excel file is empty.")

    df = df.fillna("")

    students = df.to_dict(orient="records")

    return students


def replace_text_in_paragraph(paragraph, replacements):
    """
    Replace placeholders inside a paragraph.
    """

    for old_text, new_text in replacements.items():
        if old_text in paragraph.text:
            for run in paragraph.runs:
                if old_text in run.text:
                    run.text = run.text.replace(
                        old_text,
                        str(new_text)
                    )


def replace_text_in_table(table, replacements):
    """
    Replace placeholders inside Word tables.
    """

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_text_in_paragraph(
                    paragraph,
                    replacements
                )


def generate_certificate(student):
    """
    Generate one certificate for one student.
    """

    if not TEMPLATE_FILE.exists():
        raise FileNotFoundError(
            f"Certificate template not found: {TEMPLATE_FILE}"
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document(TEMPLATE_FILE)

    replacements = {}

    for key, value in student.items():
        replacements[f"{{{{{key}}}}}"] = str(value)

    # Replace placeholders in normal paragraphs
    for paragraph in document.paragraphs:
        replace_text_in_paragraph(
            paragraph,
            replacements
        )

    # Replace placeholders inside tables
    for table in document.tables:
        replace_text_in_table(
            table,
            replacements
        )

    student_name = str(
        student.get("Name", "Student")
    ).strip()

    if not student_name:
        student_name = "Student"

    safe_name = "".join(
        character
        for character in student_name
        if character.isalnum()
        or character in (" ", "_", "-")
    ).strip()

    output_file = OUTPUT_DIR / f"{safe_name}_Certificate.docx"

    document.save(output_file)

    return str(output_file)


def generate_all_certificates():
    """
    Generate certificates for every student in the Excel file.
    """

    students = read_student_data()

    generated_files = []

    for student in students:
        output_file = generate_certificate(student)
        generated_files.append(output_file)

    return generated_files